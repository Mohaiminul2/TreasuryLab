"""
export_utils.py

Generates Excel and PDF export files summarizing a company's historical
metrics and scenario forecast. Kept separate from the Dash page so it
can also be run/tested standalone.
"""

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


def build_excel_report(ticker: str, historical_df: pd.DataFrame, forecast_df: pd.DataFrame) -> bytes:
    """Build an Excel workbook with historical metrics + scenario forecast, return as bytes."""
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Historical Metrics"
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for row in dataframe_to_rows(historical_df.round(2), index=False, header=True):
        ws1.append(row)
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws2 = wb.create_sheet("Scenario Forecast")
    for row in dataframe_to_rows(forecast_df.round(2), index=False, header=True):
        ws2.append(row)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    for ws in [ws1, ws2]:
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 30)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_pdf_report(ticker: str, historical_df: pd.DataFrame, forecast_df: pd.DataFrame) -> bytes:
    """Build a PDF summary report, return as bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"Working Capital & Cash Flow Forecast — {ticker}", styles["Title"]))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d')}", styles["Normal"]))
    elements.append(Spacer(1, 20))

    # Historical Working Capital Metrics
    elements.append(Paragraph("Historical Working Capital Metrics", styles["Heading2"]))
    hist_display = historical_df[
        [c for c in ["fiscal_year", "revenue", "ccc", "dso", "dpo", "dio", "current_ratio"] if c in historical_df.columns]
    ].round(1)

    # Handle empty dataframe case
    if hist_display.empty:
        hist_table_data = [["No historical data available"]]
    else:
        hist_table_data = [hist_display.columns.tolist()] + hist_display.values.tolist()

    hist_table = Table(hist_table_data, repeatRows=1)
    hist_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    elements.append(hist_table)
    elements.append(Spacer(1, 20))

    # Scenario Forecast (12-Month Forward)
    elements.append(Paragraph("Scenario Forecast (12-Month Forward)", styles["Heading2"]))
    fc_display = forecast_df[
        [c for c in ["scenario", "forecast_year", "revenue", "ccc", "operating_cf_estimate", "cash_flag"] if c in forecast_df.columns]
    ].round(1)

    # Handle empty dataframe case
    if fc_display.empty:
        fc_table_data = [["No forecast data available"]]
    else:
        fc_table_data = [fc_display.columns.tolist()] + fc_display.values.tolist()

    fc_table = Table(fc_table_data, repeatRows=1)
    fc_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    elements.append(fc_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()